#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""极简 Markdown 转 docx：标题、段落（支持 **粗体**）、表格、代码块、列表。用法：python3 tools/md2docx.py in.md out.docx"""
import sys, re
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

def set_font(run, size=10.5, bold=None, mono=False, color=None):
    run.font.size = Pt(size)
    name = "Consolas" if mono else "Microsoft YaHei"
    run.font.name = name; run._element.rPr.rFonts.set(qn("w:eastAsia"), name if not mono else "Microsoft YaHei")
    if bold is not None: run.bold = bold
    if color: run.font.color.rgb = RGBColor.from_string(color)

def add_inline(p, text, size=10.5, color=None):
    for part in re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", text):
        if not part: continue
        if part.startswith("**") and part.endswith("**"):
            set_font(p.add_run(part[2:-2]), size, bold=True, color=color)
        elif part.startswith("`") and part.endswith("`"):
            set_font(p.add_run(part[1:-1]), size - 0.5, mono=True, color="444444")
        else:
            set_font(p.add_run(part), size, color=color)

def shade(cell, hexcolor):
    tcPr = cell._element.get_or_add_tcPr(); sh = OxmlElement("w:shd")
    sh.set(qn("w:val"), "clear"); sh.set(qn("w:color"), "auto"); sh.set(qn("w:fill"), hexcolor); tcPr.append(sh)

def table(doc, rows):
    cells = [[c.strip() for c in r.strip().strip("|").split("|")] for r in rows if not re.match(r"^\|?\s*:?-{2,}", r.strip())]
    if not cells: return
    ncol = max(len(r) for r in cells)
    t = doc.add_table(rows=len(cells), cols=ncol); t.style = "Table Grid"
    for i, r in enumerate(cells):
        for j in range(ncol):
            c = t.cell(i, j); c.text = ""
            p = c.paragraphs[0]; p.paragraph_format.space_after = Pt(0)
            add_inline(p, r[j] if j < len(r) else "", size=9)
            if i == 0:
                shade(c, "E8E8E8")
                for run in p.runs: run.bold = True
    doc.add_paragraph().paragraph_format.space_after = Pt(2)

def main(src, out):
    doc = Document()
    for s in doc.sections:
        s.left_margin = s.right_margin = Cm(2.2); s.top_margin = s.bottom_margin = Cm(2)
    st = doc.styles["Normal"]; st.font.name = "Microsoft YaHei"; st.element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei"); st.font.size = Pt(10.5)
    lines = open(src, encoding="utf-8").read().split("\n"); i = 0
    while i < len(lines):
        l = lines[i]
        if l.startswith("```"):
            j = i + 1; buf = []
            while j < len(lines) and not lines[j].startswith("```"): buf.append(lines[j]); j += 1
            p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(6)
            set_font(p.add_run("\n".join(buf)), 8.5, mono=True, color="333333")
            i = j + 1; continue
        if l.strip().startswith("|"):
            j = i; buf = []
            while j < len(lines) and lines[j].strip().startswith("|"): buf.append(lines[j]); j += 1
            table(doc, buf); i = j; continue
        mi = re.match(r"^!\[([^\]]*)\]\(([^)]+)\)", l.strip())
        if mi:
            import os
            path = os.path.join(os.path.dirname(os.path.abspath(src)), mi.group(2))
            if os.path.exists(path):
                doc.add_picture(path, width=Cm(15.5))
                cap = doc.add_paragraph(); cap.alignment = WD_ALIGN_PARAGRAPH.CENTER; set_font(cap.add_run(mi.group(1)), 9, color="666666")
            i += 1; continue
        m = re.match(r"^(#{1,4})\s+(.*)", l)
        if m:
            lvl = len(m.group(1)); h = doc.add_heading(level=min(lvl, 4)); h.text = ""
            set_font(h.add_run(m.group(2).strip()), {1: 18, 2: 15, 3: 13, 4: 11.5}[min(lvl, 4)], bold=True, color="1F1F1F")
            i += 1; continue
        m = re.match(r"^\s*[-*]\s+(.*)", l)
        if m:
            p = doc.add_paragraph(style="List Bullet"); p.paragraph_format.space_after = Pt(2); add_inline(p, m.group(1)); i += 1; continue
        m = re.match(r"^\s*(\d+)\.\s+(.*)", l)
        if m:
            p = doc.add_paragraph(style="List Number"); p.paragraph_format.space_after = Pt(2); add_inline(p, m.group(2)); i += 1; continue
        if l.strip() == "---" or not l.strip(): i += 1; continue
        p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(6); p.paragraph_format.line_spacing = 1.25
        add_inline(p, l.strip()); i += 1
    doc.save(out); print("saved", out)

if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
